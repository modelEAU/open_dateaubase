"""Generate the two review spreadsheets for the pilEAUte metadata reconciliation.

Step 3: process units + sampling locations to create.
Step 4: per-tag -> sampling-location mapping (decoded via the loop-numbering key).

Numbering key (from domain owner):
    leading digit = process area: 0 pretreatment, 1 primary train,
                    2 pilot secondary train (Pilote), 3 copilot secondary train (Copilote)
    second digit  = reactor / tank number

Read-only against the API; writes two .xlsx files for human review.
"""

from __future__ import annotations

import os
import re
import json
import urllib.request
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

API = "http://127.0.0.1:8010/api/v1"
OUT = Path(r"C:\Users\admin_modeleau\Desktop\pileaute_reconciliation")


def _token() -> str:
    for line in Path(r"C:\source\open_dateaubase\.env.staging").read_text().splitlines():
        if line.startswith("API_SERVICE_TOKEN="):
            return line.split("=", 1)[1].strip()
    return ""


def _get(path: str) -> dict:
    req = urllib.request.Request(API + path, headers={"Authorization": f"Bearer {_token()}"})
    with urllib.request.urlopen(req, timeout=20) as r:
        return json.loads(r.read())


AREA = {"0": "Pretreatment", "1": "Primary train",
        "2": "Pilot 2ndary (Pilote)", "3": "Copilot 2ndary (Copilote)"}
INSTR = {
    "AIT": "Analyzer transmitter", "AIC": "Air-flow controller (setpoint)",
    "FIT": "Flow transmitter", "FCV": "Flow control valve", "LDO": "Luminescent DO probe",
    "LIT": "Level transmitter", "TIT": "Temperature transmitter", "TSS": "TSS sensor",
}

# ---- Process units to create (code, name, kind, parent, area) ----
PROCESS_UNITS = [
    ("PRE", "Pretreatment", "Area", "", "0 Pretreatment"),
    ("ST", "Storage tank", "Tank", "PRE", "0 Pretreatment"),
    ("PRIM", "Primary train", "Area", "", "1 Primary train"),
    ("PST", "Primary settling tank", "Clarifier", "PRIM", "1 Primary train"),
    ("PIL", "Pilot secondary train (Pilote)", "Area", "", "2 Pilot"),
    *[(f"PIL_R{i}", f"Pilote reactor {i}", "Reactor", "PIL", "2 Pilot") for i in range(1, 6)],
    ("PIL_SST", "Pilote secondary settling tank", "Clarifier", "PIL", "2 Pilot"),
    ("COP", "Copilot secondary train (Copilote)", "Area", "", "3 Copilot"),
    *[(f"COP_R{i}", f"Copilote reactor {i}", "Reactor", "COP", "3 Copilot") for i in range(1, 6)],
    ("COP_SST", "Copilote secondary settling tank", "Clarifier", "COP", "3 Copilot"),
]

# ---- Sampling locations to create (legacy_id, code, name, area, process_unit, include) ----
SAMPLING_LOCATIONS = [
    (14, "PIL_INF", "Pilote influent", "2 Pilot", "PIL", "yes"),
    (3, "PIL_R1", "Pilote reactor 1", "2 Pilot", "PIL_R1", "yes"),
    (4, "PIL_R2", "Pilote reactor 2", "2 Pilot", "PIL_R2", "yes"),
    (5, "PIL_R3", "Pilote reactor 3", "2 Pilot", "PIL_R3", "yes"),
    (6, "PIL_R4", "Pilote reactor 4", "2 Pilot", "PIL_R4", "yes"),
    (7, "PIL_R5", "Pilote reactor 5", "2 Pilot", "PIL_R5", "yes"),
    (8, "PIL_IRIN", "Pilote internal recycle IN", "2 Pilot", "PIL", "yes"),
    (9, "PIL_IROUT", "Pilote internal recycle OUT", "2 Pilot", "PIL", "yes"),
    (11, "PIL_SR", "Pilote sludge recycle", "2 Pilot", "PIL_SST", "yes"),
    (12, "PIL_SST", "Pilote secondary settling tank", "2 Pilot", "PIL_SST", "yes"),
    (10, "PIL_EFF", "Pilote effluent", "2 Pilot", "PIL_SST", "yes"),
    (13, "PIL_WW", "Pilote wet weather", "2 Pilot", "PIL", "optional"),
    (25, "COP_INF", "Copilote influent", "3 Copilot", "COP", "yes"),
    (15, "COP_R1", "Copilote reactor 1", "3 Copilot", "COP_R1", "yes"),
    (16, "COP_R2", "Copilote reactor 2", "3 Copilot", "COP_R2", "yes"),
    (17, "COP_R3", "Copilote reactor 3", "3 Copilot", "COP_R3", "yes"),
    (18, "COP_R4", "Copilote reactor 4", "3 Copilot", "COP_R4", "yes"),
    (19, "COP_R5", "Copilote reactor 5", "3 Copilot", "COP_R5", "yes"),
    (20, "COP_IRIN", "Copilote internal recycle IN", "3 Copilot", "COP", "yes"),
    (21, "COP_IROUT", "Copilote internal recycle OUT", "3 Copilot", "COP", "yes"),
    (22, "COP_SR", "Copilote sludge recycle", "3 Copilot", "COP_SST", "yes"),
    (23, "COP_SST", "Copilote secondary settling tank", "3 Copilot", "COP_SST", "yes"),
    (26, "COP_EFF", "Copilote effluent", "3 Copilot", "COP_SST", "yes"),
    (24, "COP_WW", "Copilote wet weather", "3 Copilot", "COP", "optional"),
    (1, "PST_INF", "Primary settling tank influent", "1 Primary train", "PST", "yes"),
    (2, "PST_EFF", "Primary settling tank effluent", "1 Primary train", "PST", "yes"),
    (37, "ST_INF", "Storage tank influent", "0 Pretreatment", "ST", "yes"),
    (38, "ST_EFF", "Storage tank effluent", "0 Pretreatment", "ST", "yes"),
    (33, "METEOSTFOY", "Weather station Sainte-Foy", "0 Pretreatment", "", "optional"),
    (49, "CWJB", "Federal weather station (U. Laval)", "0 Pretreatment", "", "optional"),
]


def decode(identifier: str, tag: str):
    """Return (instrument, decoded_area, unit_digit, proposed_sp, confidence, notes)."""
    src = identifier or tag
    up = src.upper()

    # monEAU equipment-backed sensors
    if identifier:
        if "INFLPC" in up or "INFL" in up:
            return ("monEAU multiprobe", "? influent", "", "PIL_INF",
                    "Low", "INFLPC1 = influent pilot channel 1? confirm train (PIL vs PST)")
        if "RODTOX" in up:
            return ("RODTOX respirometer", "? influent", "", "PIL_INF",
                    "Low", "Respirometer; confirm bypass / influent location")
        m = re.search(r"R?(\d)(\d)\d", up)  # R100/R200/R300 or pH200/pHTemp300
        if m:
            area_d, unit_d = m.group(1), m.group(2)
            area = AREA.get(area_d, "?")
            if area_d == "1":
                return ("monEAU station", area, unit_d, "PST",
                        "Low", "Primary-train monEAU station; confirm exact SP")
            train = "PIL" if area_d == "2" else ("COP" if area_d == "3" else "?")
            sp = f"{train} (reactor TBD)" if unit_d == "0" else f"{train}_R{unit_d}"
            conf = "Med" if unit_d == "0" else "High"
            note = "monEAU station, reactor digit 0 -> confirm which reactor" if unit_d == "0" else ""
            return ("monEAU station", area, unit_d, sp, conf, note)
        return ("monEAU", "?", "", "", "Low", "Unrecognized identifier")

    # PLC [PCL_001] tags
    t = re.sub(r"^\[PCL_001\]", "", tag)
    mtype = re.match(r"([A-Z]+)", t)
    instr = mtype.group(1) if mtype else "?"
    instr_name = INSTR.get(instr, instr)
    mloop = re.search(r"(\d)(\d)\d", t)  # first 3-digit loop number
    if not mloop:
        return (instr_name, "?", "", "", "Low", "No loop number parsed")
    area_d, unit_d = mloop.group(1), mloop.group(2)

    if area_d == "4":
        return (instr_name, "4 Aeration?", unit_d, "Aeration system (reactor air supply)",
                "Low", "Leading digit 4 not in key -> blower/air distribution? map to served reactor")
    area = AREA.get(area_d, "?")
    if area_d == "0":
        return (instr_name, area, unit_d, "Influent / pretreatment",
                "Med", "Pretreatment flow/level; confirm exact SP (influent vs ST)")
    if area_d == "1":
        return (instr_name, area, unit_d, "PST / primary train",
                "Med", "Primary-train instrument; confirm SP")
    train = "PIL" if area_d == "2" else "COP"
    if unit_d in "12345":
        return (instr_name, area, unit_d, f"{train}_R{unit_d}", "High", "")
    if unit_d in "67":
        return (instr_name, area, unit_d, f"{train}_SST / {train}_EFF",
                "Low", "Unit 6/7 -> settler/effluent analyzers? confirm")
    return (instr_name, area, unit_d, f"{train} (TBD)", "Low", "Confirm unit")


# ---------- styling helpers ----------
HDR = Font(bold=True, color="FFFFFF")
HDRFILL = PatternFill("solid", fgColor="305496")
CONF = {"High": PatternFill("solid", fgColor="C6EFCE"),
        "Med": PatternFill("solid", fgColor="FFEB9C"),
        "Low": PatternFill("solid", fgColor="FFC7CE")}


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HDR
        cell.fill = HDRFILL
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def autofit(ws):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 3, 60)


def main():
    OUT.mkdir(parents=True, exist_ok=True)

    # ===== Step 3 workbook =====
    wb3 = Workbook()
    ws = wb3.active
    ws.title = "ProcessUnits"
    cols = ["Tag/Code", "Name", "Kind", "Parent", "Area", "Action", "Notes"]
    ws.append(cols)
    for code, name, kind, parent, area in PROCESS_UNITS:
        ws.append([code, name, kind, parent, area, "CREATE", ""])
    style_header(ws, len(cols))
    autofit(ws)

    ws2 = wb3.create_sheet("SamplingLocations")
    cols2 = ["LegacyID", "Code", "Name", "Area", "ProcessUnit", "Create?", "Notes"]
    ws2.append(cols2)
    for lid, code, name, area, pu, inc in SAMPLING_LOCATIONS:
        ws2.append([lid, code, name, area, pu, inc, ""])
    style_header(ws2, len(cols2))
    autofit(ws2)
    p3 = OUT / "step3_process_units_and_sampling_locations.xlsx"
    wb3.save(p3)

    # ===== Step 4 workbook =====
    data = _get("/channels?page_size=1000")
    items = data.get("items", data)
    wb4 = Workbook()
    ws = wb4.active
    ws.title = "TagMapping"
    cols = ["ChannelID", "Family", "Equipment / Placeholder", "Tag", "Parameter", "Unit",
            "ValueKind", "Instrument", "DecodedArea", "UnitDigit",
            "ProposedSamplingLocation", "Confidence", "Notes"]
    ws.append(cols)

    def sort_key(c):
        return (0 if c.get("equipment_identifier") else 1, c["tag_name"])

    for c in sorted(items, key=sort_key):
        ident = c.get("equipment_identifier") or ""
        tag = c["tag_name"]
        instr, area, unit, sp, conf, note = decode(ident, tag)
        family = "monEAU (wired)" if ident else "PLC (needs placeholder)"
        placeholder = ident or ("pilEAUte_" + re.sub(r"^\[PCL_001\]", "", tag).split("_EU")[0]
                                .replace(".", "_"))
        ws.append([c["channel_id"], family, placeholder, tag,
                   c.get("parameter_name") or "", c.get("unit_name") or "",
                   c.get("value_kind_name") or "", instr, area, unit, sp, conf, note])
        ws.cell(row=ws.max_row, column=12).fill = CONF.get(conf, CONF["Low"])

    style_header(ws, len(cols))
    autofit(ws)
    p4 = OUT / "step4_tag_to_location_mapping.xlsx"
    wb4.save(p4)

    print("WROTE", p3)
    print("WROTE", p4)
    print("channels:", len(items),
          "| high:", sum(1 for c in items
                          if decode(c.get('equipment_identifier') or '', c['tag_name'])[4] == 'High'))


if __name__ == "__main__":
    main()

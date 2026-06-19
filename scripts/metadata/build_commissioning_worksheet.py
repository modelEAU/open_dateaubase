"""Step C: export the pilEAUte commissioning worksheet.

One row per signal interface (= one physical sensor). Prefilled with the interface,
DAS, the parameters it carries, and any current (placeholder) equipment. Field columns
to fill at the plant: EquipmentIdentifier, EquipmentModel, SerialNumber,
SamplingLocation, DeployFrom, Notes -- with dropdowns validated against the imported
model catalog and the Step 3 sampling-location codes.

Read-only against the API.  Output: <legacy>/commissioning_worksheet.xlsx
"""

from __future__ import annotations

import json
import urllib.request
from pathlib import Path
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

API = "http://127.0.0.1:8010/api/v1"
OUT = Path(r"C:\Users\admin_modeleau\Desktop\legacy_metadata") / "commissioning_worksheet.xlsx"
CAMPAIGN_START = "2026-01-01"


def _token() -> str:
    for line in Path(r"C:\source\open_dateaubase\.env.staging").read_text().splitlines():
        if line.startswith("API_SERVICE_TOKEN="):
            return line.split("=", 1)[1].strip()
    raise SystemExit("token not found")


TOKEN = _token()


def get(path: str):
    req = urllib.request.Request(API + path, headers={"Authorization": f"Bearer {TOKEN}"})
    with urllib.request.urlopen(req, timeout=30) as r:
        d = json.loads(r.read())
    return d.get("items", d) if isinstance(d, dict) else d


HDR = Font(bold=True, color="FFFFFF")
HDRFILL = PatternFill("solid", fgColor="305496")
LOCK = PatternFill("solid", fgColor="EDEDED")   # read-only (prefilled) columns
FILLCOL = PatternFill("solid", fgColor="FFF2CC")  # columns to fill in the field


def main():
    interfaces = get("/signal-interfaces")
    channels = get("/channels?page_size=1000")
    models = get("/equipment/models")
    sps = get("/sites/1/sampling-locations")

    # group channels by signal interface
    by_si: dict[int, list] = defaultdict(list)
    cur_equip: dict[int, str] = {}
    for c in channels:
        by_si[c["signal_interface_id"]].append(c)
        if c.get("equipment_identifier"):
            cur_equip[c["signal_interface_id"]] = c["equipment_identifier"]

    model_strings = sorted(
        {f"{(m.get('manufacturer') or '?')} - {m.get('equipment_model') or '?'}" for m in models}
    )
    sp_codes = [s["name"] for s in sps]

    wb = Workbook()
    ws = wb.active
    ws.title = "Commissioning"
    cols = ["SignalInterfaceID", "SignalInterface", "DAS", "Parameters", "CurrentEquipment",
            "EquipmentIdentifier", "EquipmentModel", "SerialNumber", "SamplingLocation",
            "DeployFrom", "Notes"]
    ws.append(cols)

    rows = sorted(interfaces, key=lambda s: (s["name"].startswith("[PCL"), s["name"]))
    for si in rows:
        sid = si["signal_interface_id"]
        chans = by_si.get(sid, [])
        params = ", ".join(sorted({c.get("parameter_name") or c["tag_name"] for c in chans}))
        current = cur_equip.get(sid, "")
        ws.append([sid, si["name"], si.get("das_name"), params, current,
                   current,            # EquipmentIdentifier default = current (editable)
                   "", "", "",         # model / serial / sampling location -> fill
                   CAMPAIGN_START, ""])

    # styling
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HDR
        cell.fill = HDRFILL
        cell.alignment = Alignment(vertical="center")
    n = ws.max_row
    for col_idx in (1, 2, 3, 4, 5):       # prefilled / read-only
        for r in range(2, n + 1):
            ws.cell(row=r, column=col_idx).fill = LOCK
    for col_idx in (6, 7, 8, 9, 10, 11):  # to fill
        for r in range(2, n + 1):
            ws.cell(row=r, column=col_idx).fill = FILLCOL
    ws.freeze_panes = "A2"

    # Reference sheet for dropdowns
    ref = wb.create_sheet("Reference")
    ref["A1"] = "EquipmentModel (Manufacturer - Model)"
    ref["A1"].font = Font(bold=True)
    for i, m in enumerate(model_strings, start=2):
        ref.cell(row=i, column=1, value=m)
    ref["C1"] = "SamplingLocation code"
    ref["D1"] = "Description"
    ref["C1"].font = ref["D1"].font = Font(bold=True)
    for i, s in enumerate(sps, start=2):
        ref.cell(row=i, column=3, value=s["name"])
        ref.cell(row=i, column=4, value=s.get("description"))

    # data validation dropdowns
    dv_model = DataValidation(
        type="list", formula1=f"=Reference!$A$2:$A${len(model_strings) + 1}", allow_blank=True)
    dv_sp = DataValidation(
        type="list", formula1=f"=Reference!$C$2:$C${len(sp_codes) + 1}", allow_blank=True)
    ws.add_data_validation(dv_model)
    ws.add_data_validation(dv_sp)
    dv_model.add(f"G2:G{n}")  # EquipmentModel
    dv_sp.add(f"I2:I{n}")     # SamplingLocation

    # column widths
    widths = [16, 42, 16, 40, 22, 22, 36, 16, 18, 12, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ref.column_dimensions["A"].width = 40
    ref.column_dimensions["C"].width = 16
    ref.column_dimensions["D"].width = 44

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("WROTE", OUT)
    print(f"rows: {n - 1} signal interfaces | models in dropdown: {len(model_strings)} | "
          f"sampling locations: {len(sp_codes)}")


if __name__ == "__main__":
    main()
